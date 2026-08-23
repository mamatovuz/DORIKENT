"""Excel bilan ishlash: shablon yaratish va yuklangan faylni o'qish."""
import io
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


HEADERS = ["Savol", "To'g'ri javob", "Noto'g'ri 1", "Noto'g'ri 2", "Noto'g'ri 3"]

SAMPLE_ROWS = [
    ["Mijoz bilan ishlashda eng muhim qoida qaysi?",
     "Mijozni tinglash va hurmat bilan muomala qilish",
     "Mijozga tezroq javob berish", "Mijoz bilan bahslashmaslik",
     "Muammoni boshqa xodimga berish"],
    ["Mijoz shikoyat qilsa nima qilish kerak?",
     "Diqqat bilan tinglab, muammoni hal qilishga harakat qilish",
     "", "", ""],  # noto'g'rilar bo'sh -> AI/offline yaratadi
    ["Ish vaqtida telefon jiringlasa?",
     "Xushmuomalalik bilan javob berish", "", "", ""],
]


def build_template() -> io.BytesIO:
    """Namuna (shablon) Excel faylini xotirada yaratadi."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Savollar"

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, row in enumerate(SAMPLE_ROWS, start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    # Ustunlar kengligi
    widths = [55, 45, 30, 30, 30]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Izoh varag'i
    ws2 = wb.create_sheet("Yo'riqnoma")
    notes = [
        "YO'RIQNOMA",
        "",
        "1-ustun: Savol matni (majburiy).",
        "2-ustun: To'g'ri javob (majburiy).",
        "3,4,5-ustun: Noto'g'ri variantlar (ixtiyoriy).",
        "",
        "Agar noto'g'ri variantlarni yozmasangiz,",
        "bot ularni AI (Gemini/ChatGPT) yordamida avtomatik yaratadi.",
        "AI ulanmagan bo'lsa — AI-siz generator ishlaydi.",
        "",
        "Birinchi qatorni (sarlavhani) o'chirmang.",
        "Har bir qator = bitta savol.",
    ]
    for i, line in enumerate(notes, start=1):
        cell = ws2.cell(row=i, column=1, value=line)
        if i == 1:
            cell.font = Font(bold=True, size=14)
    ws2.column_dimensions["A"].width = 60

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class ParsedQuestion:
    def __init__(self, text: str, correct: str, wrongs: list[str]):
        self.text = text
        self.correct = correct
        self.wrongs = [w for w in wrongs if w and str(w).strip()]

    @property
    def needs_ai(self) -> bool:
        return len(self.wrongs) < 3


def parse_excel(data: bytes) -> tuple[list[ParsedQuestion], list[str]]:
    """
    Yuklangan Excel baytlarini o'qiydi.
    Qaytadi: (savollar ro'yxati, ogohlantirishlar ro'yxati)
    """
    warnings: list[str] = []
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb["Savollar"] if "Savollar" in wb.sheetnames else wb.worksheets[0]

    questions: list[ParsedQuestion] = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None:
            continue
        cells = list(row) + [None] * (5 - len(row)) if len(row) < 5 else list(row)
        q_text = (str(cells[0]).strip() if cells[0] is not None else "")
        correct = (str(cells[1]).strip() if cells[1] is not None else "")
        if not q_text and not correct:
            continue  # bo'sh qator
        if not q_text:
            warnings.append(f"{idx}-qator: savol matni yo'q — o'tkazib yuborildi.")
            continue
        if not correct:
            warnings.append(f"{idx}-qator: to'g'ri javob yo'q — o'tkazib yuborildi.")
            continue
        wrongs = [str(cells[i]).strip() for i in (2, 3, 4)
                  if cells[i] is not None and str(cells[i]).strip()]
        questions.append(ParsedQuestion(q_text, correct, wrongs))

    return questions, warnings


def build_results_excel(rows) -> io.BytesIO:
    """Natijalar ro'yxatini (aiosqlite Row lar) Excel qilib qaytaradi."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Natijalar"
    headers = ["#", "Xodim", "Username", "Test", "Savollar", "To'g'ri",
               "Noto'g'ri", "Natija %", "Holat", "Sana"]
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = header_fill
        c.font = header_font
    for i, r in enumerate(rows, start=1):
        ws.append([
            i,
            r["full_name"] or "-",
            f"@{r['username']}" if r["username"] else "-",
            r["title"],
            r["total"],
            r["correct"],
            r["wrong"],
            round(r["percent"], 1),
            "O'tdi" if r["passed"] else "O'tmadi",
            r["created_at"],
        ])
    widths = [5, 22, 18, 25, 10, 8, 10, 10, 10, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
